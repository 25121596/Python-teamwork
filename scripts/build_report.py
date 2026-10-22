"""Build the teacher-template report from verified public experiment evidence.

Document authoring uses python-docx in the artifact runtime. The scientific
runtime is separate. Teacher input is kept byte-identical under docs/references.
"""
from __future__ import annotations

import csv
import json
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables'
TEMPLATE = ROOT / 'docs/references/项目报告模板.docx'


def load(name):
    return json.loads((ROOT / name).read_text(encoding='utf-8'))


def build():
    eda = load('artifacts/audit/eda_summary.json')
    selection = load('artifacts/metrics/model_selection.json')
    ablation = load('artifacts/metrics/cleaning_ablation.json')
    final = load('artifacts/metrics/final_evaluation.json')
    final_audit = load('artifacts/audit/final_evaluation_validation.json')
    review = load('artifacts/metrics/error_review_assisted.json')
    errors = load('artifacts/metrics/error_analysis.json')
    if not final['passed'] or not final_audit['passed'] or review['ai_assisted_review_status'] != 'complete':
        raise ValueError('Verified real evidence and assisted review are required.')
    m = final['metrics']; lr = m['lr']; nb = m['nb']
    cmp = {x['representation']: x for x in ablation['comparisons']}
    pct = lambda x: f'{100*x:.2f}%'
    cm = lr['confusion_matrix']
    representation = '去除 HTML 标签的可读文本' if final['representation'] == 'text_clean' else '保留 HTML 标签文本'
    doc = Document(TEMPLATE)
    original = [deepcopy(x) for x in doc._element.body]
    for x in list(doc._element.body):
        if x.tag != qn('w:sectPr'): doc._element.body.remove(x)
    sec = doc.sections[0]
    width = (sec.page_width - sec.left_margin - sec.right_margin) / 914400
    for name, size, font in [('Normal',12,'宋体'),('Title',36,'黑体'),('Heading 1',14,'黑体'),('Heading 2',12,'黑体'),('Caption',10,'宋体')]:
        if name not in doc.styles: doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        s=doc.styles[name]; s.font.name=font; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
        rf=s.element.get_or_add_rPr().get_or_add_rFonts()
        for key in list(rf.attrib):
            if 'theme' in key.lower(): del rf.attrib[key]
        rf.set(qn('w:eastAsia'),font)
        f=s.paragraph_format; f.space_before=Pt(0); f.space_after=Pt(6); f.line_spacing=Pt(22 if size<=14 else 45)
        f.widow_control=True; f.keep_with_next=name in ('Heading 1','Heading 2','Title')
        snap=OxmlElement('w:snapToGrid'); snap.set(qn('w:val'),'0'); s.element.get_or_add_pPr().append(snap)
    for name in ('Heading 1','Heading 2','Title'): doc.styles[name].font.bold=True
    for s in doc.styles:
        for e in list(s.element.iter(qn('w:pBdr'))): e.getparent().remove(e)
    for name in ('Heading 1','Heading 2'):
        doc.styles[name].paragraph_format.space_before=Pt(8)
    # Preserve the teacher's A4 section, cover labels and typography. Compress
    # surplus blank paragraphs so the last cover field stays on the cover.
    cover_indices=[0,1,3,4,5,7,8,18,19,20,21,22]
    for idx in cover_indices: doc._element.body.insert(len(doc._element.body)-1,original[idx])
    cp=doc.paragraphs
    for para in cp:
        para.paragraph_format.keep_with_next=False
        para.paragraph_format.widow_control=True
    cp[0].text='评分：_______'; cp[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
    cp[1].paragraph_format.space_after=Pt(22)
    cp[2].style=doc.styles['Title']; cp[2].paragraph_format.space_after=Pt(25)
    cp[3].style=doc.styles['Title']; cp[3].paragraph_format.space_after=Pt(20)
    cp[5].text='中文垃圾邮件分类'; cp[5].style=doc.styles['Title']; cp[5].alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in cp[5].runs:
        run.font.name='黑体'; run.font.size=Pt(22); run.bold=True
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'黑体')
    cp[6].paragraph_format.space_after=Pt(35)
    for para,label in zip(cp[7:],['学    院','专    业','学    号','学生姓名','课    程']):
        para.clear(); para.paragraph_format.first_line_indent=Inches(.45)
        para.paragraph_format.line_spacing=Pt(28); para.paragraph_format.space_after=Pt(5)
        r=para.add_run(label+'  '); r.font.size=Pt(16); r.bold=True
        r=para.add_run('Python数据分析编程' if label=='课    程' else '____________________')
        r.underline=True; r.font.size=Pt(14)
    p=doc.add_paragraph('实验完成版  成员信息待本人补齐',style='Caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer=sec.footer.paragraphs[0]; footer.clear(); footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ').font.size=Pt(9)
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    footer.add_run(' 页').font.size=Pt(9)
    md=['# 中文垃圾邮件分类项目报告','', '教师原模板版，真实实验完成，成员信息与实际贡献待本人确认。','']

    next_page = False
    def page():
        nonlocal next_page
        next_page = True
    def p(text,style=None):
        para=doc.add_paragraph(text,style or 'Normal')
        if style is None: para.paragraph_format.first_line_indent=Pt(24)
        md.extend([text.replace('<','&lt;').replace('>','&gt;'),'']); return para
    def h(text,level=1):
        nonlocal next_page
        para=doc.add_heading(text,level)
        if next_page:
            para.paragraph_format.page_break_before=True
            next_page=False
        md.extend(['#'*(level+1)+' '+text,'']); return para
    def table(headers,rows,widths=None,size=10.5):
        widths=widths or [width/len(headers)]*len(headers)
        factor=width/sum(widths); widths=[w*factor for w in widths]
        t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for c,w in zip(t.columns,widths): c.width=Inches(w)
        for c,v in zip(t.rows[0].cells,headers): c.text=str(v)
        for values in rows:
            for c,v in zip(t.add_row().cells,values): c.text=str(v)
        repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for ri,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for c,w in zip(row.cells,widths):
                c.width=Inches(w); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    e=OxmlElement('w:'+edge)
                    for k,v in [('val','single'),('sz','4'),('color','BBBBBB')]: e.set(qn('w:'+k),v)
                    borders.append(e)
                pr.append(borders)
                if ri==0:
                    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EFEFEF'); pr.append(shade)
                for para in c.paragraphs:
                    pf=para.paragraph_format; pf.line_spacing=Pt(16); pf.space_after=Pt(4); pf.space_before=Pt(4); pf.first_line_indent=Pt(0)
                    for r in para.runs: r.font.size=Pt(size); r.bold=ri==0
        md.extend(['| '+' | '.join(headers)+' |','|'+'---|'*len(headers)])
        md.extend('| '+' | '.join(map(str,row))+' |' for row in rows); md.append('')
        return t
    def fig(name,caption,w=None):
        para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.line_spacing=1.0; para.paragraph_format.keep_with_next=True
        para.add_run().add_picture(str(ROOT/'artifacts/figures'/name),width=Inches(w or width))
        p(caption,'Caption'); md.extend([f'![{caption}](../artifacts/figures/{name})',''])
    def code(text):
        para=p(text); para.paragraph_format.first_line_indent=Pt(0); para.paragraph_format.line_spacing=Pt(13)
        for r in para.runs: r.font.name='Consolas'; r.font.size=Pt(9)
        md[-2:] = ['```python' if not text.startswith('python ') else '```bash',text,'```','']

    page()
    h('目录')
    toc=[('1 引言',3),('2 数据介绍',3),('3 数据清洗与预处理',4),('4 探索性分析',5),('5 特征工程与模型构建',6),('6 评估与结果解释',8),('7 总结与结论',10),('8 小组成员分工与个人承担',10),('9 附录',11)]
    for title,num in toc:
        para=doc.add_paragraph(f'{title}\t{num}')
        para.paragraph_format.tab_stops.add_tab_stop(Inches(width-.05),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        para.paragraph_format.space_after=Pt(1)
    h('摘要')
    p(f'本项目处理课程提供的 TREC 2006 中文邮件语料，研究邮件解析、编码恢复、HTML 清洗与文本分类。64,620 封邮件经清洗和精确去重后保留 33,988 封代表，按内容分组冻结训练、验证、测试集合。采用 jieba 分词和训练侧 50,000 维 TF-IDF，比较多数类、朴素贝叶斯与逻辑回归。通过验证集选参及 HTML 对照后，锁定 LR C=2 和{representation}。在 5,097 封最终测试邮件上，LR 准确率为 {pct(lr["accuracy"])}，macro-F1 为 {pct(lr["macro_f1"])}，正常邮件误报 {cm["fp"]} 封，垃圾邮件漏报 {cm["fn"]} 封。结果支持其在本历史语料上优于本轮 NB 基线，但不能直接外推实际收件箱。结合真实错分的 AI 辅助语义复核，报告提出复核队列和数据质量改进建议。')
    p('关键词：中文垃圾邮件；邮件解析；TF-IDF；朴素贝叶斯；逻辑回归','Caption')

    page(); h('1 引言')
    p('垃圾邮件识别需要同时控制误拦正常邮件和漏掉垃圾邮件。本项目可靠提取邮件文本，公平比较两种分类方法，检验 HTML 清洗影响，并解释真实错误与关键词。除准确率外，重点关注正常邮件误报率与垃圾邮件召回率。')
    p('固定开发样本用于解析调试；训练集学习词表、IDF 与模型参数；验证集用于有限选参和清洗对照；锁定后才计算最终测试成绩。每阶段保存来源指纹与检查记录。')
    h('2 数据介绍')
    table(['范围','正常邮件','垃圾邮件','合计'],[['主索引','21,766','42,854','64,620'],['技术可用','21,766','42,854','64,620'],['去重代表','21,411','12,577','33,988']])
    p('数据由课程附件提供，入口为 trec06c/full/index。每行给出 ham 或 spam 标签及邮件相对路径，标签映射为 ham=0、spam=1。delay/index 涉及同一批邮件，不能视作新增样本或独立测试集。原始邮件为 RFC822 格式，正文多含 HTML，字符编码以 GBK/GB2312 为主。')
    table(['字段','用途'],[['mail_id','由相对路径确定的稳定 ID，仅用于对账'],['label / label_name','二元真值和类别名称，不进入特征'],['subject_clean / body_clean','规范化后的主题和正文'],['text_clean / HTML 对照','去标签文本与保留 HTML 的对照表示'],['parse_status / 质量标记','解析状态、HTML 与编码警告等质量信息']],[1.9,4])
    p('先固定 ham、spam 各 200 封开发样本用于解析器调试；正式分析采用全量语料。原始邮件、逐封文本、完整词表、模型文件和预测明细只保存在本地忽略目录。公开仓库保留代码、聚合结果、固定划分元数据和阶段说明。')

    page(); h('3 数据清洗与预处理')
    h('3.1 编码与结构恢复',2)
    p('二进制读取邮件后解析头部与 MIME。先解 base64、quoted-printable 等传输编码，再按声明字符集与明确候选回退解码。multipart/alternative 选择一个有效正文表示，避免纯文本和 HTML 重复拼接；related 保留根内容，跳过二进制附件与转发的整封子邮件。损坏边界、回退动作和质量警告分别记录。')
    p('HTML 去掉 script/style 与标签，解码实体并保留可读文字和锚文本。随后进行 Unicode NFKC、英文小写、控制字符与空白清理，以及常见联系字段归一化。主题与正文分别清洗后组合。空主题或正文不填造内容，只有主题可用时保留并标记；不可恢复样本按规则进入隔离账目。')
    table(['互斥解析状态','封数','解释'],[['ok','29,891','常规处理可用'],['recovered','33,307','记录回退或结构恢复'],['subject_only','1,422','仅主题可用'],['隔离或错误','0','本批数据的技术结果']],[1.5,.8,3.6])
    p('“技术可用”只表示获得了可供后续处理的文本，不保证语义完全正确。零替换字符也不能证明不存在乱码。异常长度只作诊断，不仅凭长度删掉邮件。常见联系字段掩码并非全面匿名化，因此真实正文不公开。')
    h('3.2 清洗机制说明',2)
    p('例如人工构造的 HTML“<b>会议通知</b><script>代码</script>”应只保留“会议通知”；同一邮件的纯文本与 HTML 替代正文只选一份。这些示例用于解释规则，不充当真实误分类案例。真实清洗影响通过第 6.1 节固定条件实验衡量。')
    p('本轮附件恢复后，主索引、64,620 封清洗缓存及清洗聚合记录与历史指纹一致。开发和全量行级检查均通过。代码没有利用最终测试的分类成绩修改解析规则。')

    page(); h('3.3 去重与冻结数据划分',2)
    p('对同标签且规范化主题与正文完全相同的邮件保留一个确定代表，移除 30,632 封重复记录。相同原始字节、最终文本、有效正文或 HTML 对照文本形成关联分组，整组划分，减少精确内容跨集合泄漏。开发样本触及的 372 个分组、1,738 封代表强制进入训练侧。')
    table(['集合','正常邮件','垃圾邮件','合计'],[[label,*[f'{eda["counts"]["split_counts"][name][k]:,}' for k in ('ham','spam','total')]] for label,name in [('训练','train'),('验证','validation'),('测试','test')]])
    p('种子为 42，目标比例为 70%/15%/15%，受整组与开发约束影响实际计数略有差异。三套冻结 CSV 与原版本字节一致；ID、标签、精确内容分组和去重账本校验通过。该规则不能穷尽近重复模板。')
    h('4 探索性分析')
    fig('stage05_category_counts.png','图 1  原始邮件、清洗与去重后的类别规模',width)
    p('spam 占比由原始约 66.32% 降为去重代表中的 37.00%。ham 保留率约 98.37%，spam 约 29.35%，说明这批语料中垃圾邮件精确重复更多。评估对象是不同规范化文本，不能把去重后的类别比例当成实际邮件流分布。')

    page(); h('4.1 长度与高频词',2)
    p('训练侧 ham、spam 文本长度中位数分别为 302、499 字符；所选正文含 HTML 的比例约为 4.18%、33.88%。这些差异用于描述数据与解释工程负担，不是按测试效果筛选特征的依据。')
    fig('stage05_training_word_coverage.png','图 2  训练侧预设通用词表内的高覆盖词',width)
    p('每封邮件对每个词最多计一次，覆盖率等于包含该词的该类训练邮件数除以该类总数。“工作”“问题”等词体现日常沟通，“公司”“服务”及联系字段占位符也常见于商业邮件。公开图限定于预设通用词表且总覆盖不少于 20 封，不是完整词表的全局排名；高频不直接代表分类判别力。')
    h('5 特征工程与模型构建')
    p('jieba 采用精确模式、HMM=False，使用固定的 57 项停用词。TF-IDF 只在训练集学习词表和 IDF，min_df=2、max_features=50,000，采用次线性词频、平滑 IDF、L2 归一化和 float64 稀疏 CSR 矩阵。验证与测试仅转换，转换前后词表与 IDF 指纹不变。')
    p('TF-IDF 降低广泛出现的常见词权重，并限制单封邮件中重复词频的增长。ID、标签、文件路径、路由头和现成垃圾评分均不进入模型。空向量、矩阵行序和标签都经过核验。')

    page(); h('5.1 模型与验证选择',2)
    p('多数类由训练集确定，为所有邮件输出正常类。MultinomialNB 用平滑的类别词分布分类，比较 alpha=0.1、0.5、1；LogisticRegression 联合学习词权重，比较 C=0.5、1、2，固定 liblinear、max_iter=2000、random_state=42，并检查收敛。两种方法使用相同训练样本、标签与特征规则。')
    p('预设选择顺序为 macro-F1 降序、正常邮件误报率升序、spam 召回率降序，再按候选名稳定打破平局。使用未四舍五入的数值，排除多数类及未收敛候选。不调阈值，也不在选择后用训练加验证重新拟合。')
    headers=['验证模型','准确率','精确率','召回率','macro-F1','误报率']
    labels={'majority':'多数类','nb_alpha_1':'NB α=1','nb_alpha_0.1':'NB α=0.1','nb_alpha_0.5':'NB α=0.5','lr_C_1':'LR C=1','lr_C_0.5':'LR C=0.5','lr_C_2':'LR C=2'}
    table(headers,[[labels[x['candidate_id']],*[pct(x['metrics'][k]) for k in ('accuracy','spam_precision','spam_recall','macro_f1','ham_false_positive_rate')]] for x in selection['experiments']],[1.1,.9,.9,.9,.95,.9],10)
    p('上表只包含 5,099 封验证邮件的结果，精确率与召回率以 spam 为正类。Linux 复现中两个非最优 NB 配置的误报各差 1 封；最佳 NB 与 LR 指标不变，分别选 alpha=0.1 和 C=2。LR C=2 验证误报 11 封、漏报 47 封，按预设规则进入清洗对照。')
    h('5.2 指标定义',2)
    p('FP 是误拦正常邮件，FN 是漏掉垃圾邮件。准确率为 (TP+TN)/N；spam 精确率为 TP/(TP+FP)，召回率为 TP/(TP+FN)；F1 为精确率与召回率的调和平均，macro-F1 为 ham 与 spam 两类 F1 的均值；正常邮件误报率为 FP/(TN+FP)。分母为零时按预定规则记零。')
    p('多数类验证准确率虽达 62.99%，却漏掉全部 1,887 封垃圾邮件，说明只看准确率会掩盖失败。最终泛化评估须使用锁定后的测试集，不能把验证选参成绩当作最终成绩。')

    page(); h('6 评估与结果解释')
    h('6.1 HTML 清洗影响对照',2)
    table(['LR C=2 输入','准确率','macro-F1','误报率'],[[label,*[pct(cmp[key]['metrics'][k]) for k in ('accuracy','macro_f1','ham_false_positive_rate')]] for key,label in [('text_clean','去标签'),('text_unstripped_html','保留 HTML')]], [2,1.2,1.3,1.4])
    delta=100*(cmp['text_clean']['metrics']['macro_f1']-cmp['text_unstripped_html']['metrics']['macro_f1'])
    p(f'固定训练/验证 ID、标签、分词和 LR C=2，仅改变 HTML 表示，各自在训练集拟合 TF-IDF，独立核验通过。去标签 macro-F1 高 {delta:.3f} 个百分点，按预设规则选{representation}。实际改变输入的训练/验证样本仅 156/36 封，净少漏报 1 封，不能宣称普遍或显著提升。')
    h('6.2 锁定后的最终测试',2)
    table(['测试模型','准确率','精确率','召回率','macro-F1','误报率'],[[label,*[pct(m[key][k]) for k in ('accuracy','spam_precision','spam_recall','macro_f1','ham_false_positive_rate')]] for key,label in [('majority','多数类'),('nb','NB α=0.1'),('lr','LR C=2')]], [1.1,.9,.9,.9,.95,.9],10)
    p(f'测试前锁定方案、参数与来源，最终测试含正常 3,211 封、垃圾 1,886 封。LR 正常/垃圾类 F1 为 {pct(lr["ham_f1"])}/{pct(lr["spam_f1"])}，误报 {cm["fp"]}、漏报 {cm["fn"]}；NB 误报 {nb["confusion_matrix"]["fp"]}、漏报 {nb["confusion_matrix"]["fn"]}。独立重算与重载预测一致。')
    para=doc.add_paragraph(); para.paragraph_format.line_spacing=1.0; para.alignment=WD_ALIGN_PARAGRAPH.CENTER; para.paragraph_format.keep_with_next=True
    for name in ('nb','lr'): para.add_run().add_picture(str(ROOT/f'artifacts/figures/stage08_{name}_confusion.png'),width=Inches(width/2-.28))
    p('图 3  最终测试混淆矩阵，左 NB，右 LR。行是真实类别，列是预测类别。','Caption')
    md.extend(['![NB 混淆矩阵](../artifacts/figures/stage08_nb_confusion.png)','![LR 混淆矩阵](../artifacts/figures/stage08_lr_confusion.png)',''])

    page(); h('6.3 易错样本与模型关联词',2)
    p(f'NB、LR 各按分数绝对值从高到低选择最多 15 个 FP 和 15 个 FN，不足则全部取出，共 {review["case_count"]} 个模型错误实例。正文只在本地阅读，公开内容为去除身份、联系方式和原句的语义概述。本轮完成 AI 辅助逐例核查；长邮件检查首、中、尾选段，学生全文人工签认仍待完成。置信度排序样本并非随机抽样，以下类别不能解释为全部错误的发生比例。')
    rows=[[x['case_id'],x['safe_paraphrase']] for x in review['report_examples']]
    table(['实例','语义观察'],rows,[1,4.9],10.5)
    p(review['interpretation'])
    fig('stage08_lr_public_terms.png','图 4  预设公开词表内的 LR 正负关联词',width*.83)
    p('系数正值使线性得分更偏向 spam，负值更偏向 ham。权重反映训练分布下的条件关联，不证明因果关系，也不能把某个词等同于垃圾邮件。高频覆盖率和系数回答不同问题；两张词图都仅展示固定通用词表内的结果。')

    page(); h('7 总结与结论')
    p(f'本项目完成原始邮件解析、精确去重、冻结划分、训练侧特征、两模型验证比较和清洗对照。锁定后的 LR C=2 在最终测试中达到 {pct(lr["accuracy"])} 准确率与 {pct(lr["macro_f1"])} macro-F1，在本轮设置下优于 NB。正常邮件仍有 {cm["fp"]} 封被误拦，不能据高准确率直接启用永久删除。')
    p('建议优先将模型输出用于可恢复的垃圾箱或复核队列，为误拦保留申诉和恢复入口；对短文本、仅主题邮件及内容标签可能不一致的实例进行复核。若未来改阈值、增加字符特征或改进解析器，应重新设计训练和验证方案，并使用新的未见测试集，不能继续利用本次测试挑选改动。')
    p('局限包括语料年代较早、精确去重后类别比例变化、可能残留近重复模板、编码恢复不保证语义无误，以及有限验证选参。当前结果不代表实际收件箱的错误率，也未检验跨时间、跨来源稳定性。AI 辅助错误解释是待成员确认的观察，不是对原标签的重新标注。')
    h('8 小组成员分工与个人承担')
    table(['成员姓名','学号','承担章节或模块','具体工作与提交'],[['待本人填写','待本人填写','按实际工作填写','对应真实 Git 提交'],['待本人填写','待本人填写','按实际工作填写','对应真实 Git 提交'],['待本人填写','待本人填写','按实际工作填写','对应真实 Git 提交']], [.9,1,1.6,2.4],10.5)
    p('成员一：待本人说明负责部分的分析思路、方法步骤、遇到的问题、解决方式与对应提交。','Caption')
    p('成员二：待本人说明负责部分的分析思路、方法步骤、遇到的问题、解决方式与对应提交。','Caption')
    p('成员三：待本人说明负责部分的分析思路、方法步骤、遇到的问题、解决方式与对应提交。','Caption')
    p('本轮续作使用 AI 辅助编程、实验核验与材料整理；不把辅助提交冒充三位学生的个人贡献。组号、组长、学院、专业与成员信息仍需确认。提交课程终稿前，成员应完成真实错分复核、复现或讲解，并据实际记录填写本节。','Caption')

    page(); h('9 附录')
    h('9.1 仓库与运行方法',2)
    p('仓库：https://github.com/25121596/Python-teamwork\n续作分支：complete-project-stages。每个实际阶段配有 docs/stages Markdown 和独立 Git 提交。当前 GitHub 应用写入返回 403，续作尚未上传成功；交付 Git bundle 保留完整本地历史。','Caption')
    p('本轮使用 Python 3.12.14，依赖见 requirements.txt；原语料放在仓库外，以 config.local.json 配置。新机器按 README 从测试前版本 b4a5299 建立独立 worktree，先完成索引、清洗、冻结划分及模型训练验收，再按下列顺序运行；每步失败即停止。如实报告跨平台差异，不覆盖已锁定发布证据。')
    code('python scripts/compare_cleaning.py\npython scripts/check_ablation.py\npython scripts/evaluate_final.py\npython scripts/check_final_results.py')
    h('9.2 核心代码摘录',2)
    p('节选自 scripts/evaluate_final.py：先验证锁，再读取测试；省略重复运行检查与访问记录写入。')
    code('lock = verify_lock(root)\nlock_sha = file_sha256(root / "artifacts/metrics/final_lock.json")\nrequire_ablation_validation(root, lock_sha)\nrows = read_test_rows(root, lock)\nrecords = read_records(root, {"test": rows},\n    lock["cleaning_cache_sha256"])["test"]')
    h('9.3 证据与提交事项',2)
    p('原始要求与空白模板：docs/references/；解析、划分和模型核验：artifacts/audit/；验证、对照和最终测试表：artifacts/metrics/；公开图：artifacts/figures/。正文报告不包含真实邮件全文、逐封预测和完整词表。')
    p('PPT 与约 5 分钟真实讲解视频需在 2026 年 10 月 31 日以前提交学习通，命名为“组号_6_组长姓名.pptx/.mp4”。Word 在 11 月 8 日晚 12 点前提交学习通，纸质版交未来技术学院 418。视频录制、成员信息及课程平台提交仍由小组完成。')
    OUT.mkdir(exist_ok=True)
    doc.save(OUT/'项目报告.docx')
    (OUT/'项目报告.md').write_text('\n'.join(md)+'\n',encoding='utf-8')
    print(json.dumps({'output':str(OUT/'项目报告.docx'),'template':str(TEMPLATE),'status':'real experiments complete; member fields pending'},ensure_ascii=False))


if __name__ == '__main__': build()
